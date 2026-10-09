"""Prepare the declared offline Retail slice once from the official archive.

Run: python recommender-retail-data.py --archive retail-original.zip --output retail-data
Requires openpyxl. No fit. Each role has its own file so the study can postpone
parsing February events until after both prediction/selection locks.
"""
import argparse
from collections import Counter, defaultdict
import csv
from datetime import datetime
import hashlib
import io
import json
from pathlib import Path
import re
import zipfile


ARCHIVE_SHA256 = 'f5385cbb54bbebf7196389109c6b0621faab0c304e3702548165e71c84aede8b'
MEMBER_SHA256 = '43465a06f2ccf7c8b5bd2892bc7defb52f97487934fe93b16ae4c3936424676d'
BOUNDARIES = {'train': ('2010-12-01', '2011-01-01'),
              'validation': ('2011-01-01', '2011-02-01'),
              'test': ('2011-02-01', '2011-03-01')}
EXPECTED_FILTERS = {'allRows': 541909, 'excludedMissingCustomer': 135080,
                    'excludedCancellation': 8905, 'excludedNonpositiveQuantity': 0,
                    'excludedNonpositivePrice': 40, 'excludedNonproductCode': 1838,
                    'retainedProductRowsAllDates': 396046}
EXPECTED_ROLES = {
    'train': {'occurrences': 1930, 'pairs': 1725, 'purchaseUsers': 200, 'truthUsers': 200,
              'truthPairs': 1725, 'outsideTruthPairs': 5181},
    'validation': {'occurrences': 533, 'pairs': 497, 'purchaseUsers': 72, 'truthUsers': 68,
                   'truthPairs': 317, 'outsideTruthPairs': 1757},
    'test': {'occurrences': 410, 'pairs': 384, 'purchaseUsers': 60, 'truthUsers': 48,
             'truthPairs': 171, 'outsideTruthPairs': 1185},
}


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def write_json_exclusive(path, value):
    with Path(path).open('x', encoding='utf-8', newline='\n') as stream:
        json.dump(value, stream, indent=2, allow_nan=False)
        stream.write('\n')


def original_occurrences(archive_path):
    from openpyxl import load_workbook

    archive = Path(archive_path)
    if digest(archive) != ARCHIVE_SHA256:
        raise ValueError('official retained archive identity differs')
    with zipfile.ZipFile(archive) as zipped:
        names = [name for name in zipped.namelist() if name.lower().endswith('.xlsx')]
        if len(names) != 1:
            raise ValueError('one original workbook member is required')
        body = zipped.read(names[0])
    if hashlib.sha256(body).hexdigest() != MEMBER_SHA256:
        raise ValueError('official workbook member identity differs')
    book = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    counts = Counter({name: 0 for name in EXPECTED_FILTERS})
    events = set()
    try:
        rows = book.active.iter_rows(values_only=True)
        header = list(next(rows))
        expected_header = ['InvoiceNo', 'StockCode', 'Description', 'Quantity',
                           'InvoiceDate', 'UnitPrice', 'CustomerID', 'Country']
        if header != expected_header:
            raise ValueError('original workbook column schema differs')
        for invoice, item, _, quantity, date, price, customer, _ in rows:
            counts['allRows'] += 1
            if customer is None:
                counts['excludedMissingCustomer'] += 1
                continue
            if str(invoice).upper().startswith('C'):
                counts['excludedCancellation'] += 1
                continue
            if quantity is None or quantity <= 0:
                counts['excludedNonpositiveQuantity'] += 1
                continue
            if price is None or price <= 0:
                counts['excludedNonpositivePrice'] += 1
                continue
            stock = str(item)
            if not re.fullmatch(r'\d{5}[A-Za-z]?', stock):
                counts['excludedNonproductCode'] += 1
                continue
            counts['retainedProductRowsAllDates'] += 1
            if not isinstance(date, datetime) or date.tzinfo is not None:
                raise ValueError('original source-naive calendar datetime required')
            if int(customer) != customer:
                raise ValueError('CustomerID must have its canonical integer spelling')
            if datetime(2010, 12, 1) <= date < datetime(2011, 3, 1):
                events.add((str(int(customer)), stock, str(invoice), date.isoformat()))
    finally:
        book.close()
    if dict(counts) != EXPECTED_FILTERS or len(events) != 65248:
        raise ValueError('complete disjoint filters/occurrence count differ from the declared audit')
    return events, dict(counts), names[0]


def select_cohort(events):
    december = [row for row in events if row[3] < '2011-01-01']
    frequency = Counter(row[1] for row in december)
    catalogue = sorted(frequency, key=lambda stock: (-frequency[stock], stock))[:100]
    catalogue_set = set(catalogue)
    history = defaultdict(set)
    for customer, stock, _, _ in december:
        if stock in catalogue_set:
            history[customer].add(stock)
    eligible = [customer for customer, row in history.items() if len(row) >= 3]
    customers = sorted(eligible, key=lambda customer: (hashlib.sha256(customer.encode()).hexdigest(), customer))[:200]
    if len(catalogue) != 100 or len(eligible) != 615 or len(customers) != 200:
        raise ValueError('train-only catalogue/customer selection differs')
    return catalogue, customers


def prepare(archive_path, output_path):
    output = Path(output_path)
    output.mkdir(parents=True, exist_ok=False)
    write_json_exclusive(output/'preparation-start.json',
                         {'sourceSHA256': digest(__file__), 'archiveSHA256': digest(archive_path),
                          'fits': 0, 'purpose': 'New declared offline derivation; never overwrite after failure'})
    events, filters, member = original_occurrences(archive_path)
    catalogue, customers = select_cohort(events)
    selected = set(customers)
    selected_events = sorted((row for row in events if row[0] in selected),
                             key=lambda row: (row[3], row[0], row[1], row[2]))
    user_map = {raw: f'U{index:03d}' for index, raw in enumerate(customers)}
    item_map = {raw: f'I{index:03d}' for index, raw in enumerate(catalogue)}
    outside = sorted({row[1] for row in selected_events}-set(catalogue))
    item_map.update({raw: f'J{index:04d}' for index, raw in enumerate(outside)})
    lex_order = {raw: index for index, raw in enumerate(sorted(catalogue))}
    write_json_exclusive(output/'private-original-id-map.json',
                         {'userLocalToOriginal': {local: raw for raw, local in user_map.items()},
                          'itemLocalToOriginal': {local: raw for raw, local in item_map.items()},
                          'disposition': 'Private transformation evidence; local IDs do not imply anonymization'})
    past = defaultdict(set)
    roles = {}
    for role, (start, end) in BOUNDARIES.items():
        period = [row for row in selected_events if start <= row[3] < end]
        inside = [row for row in period if row[1] in set(catalogue)]
        truth, outside_truth = defaultdict(set), defaultdict(set)
        for customer, stock, _, _ in period:
            if stock not in past[customer]:
                (truth if stock in set(catalogue) else outside_truth)[customer].add(stock)
        counts = {'occurrences': len(inside), 'pairs': len({row[:2] for row in inside}),
                  'purchaseUsers': len({row[0] for row in inside}),
                  'truthUsers': sum(bool(truth[user]) for user in customers),
                  'truthPairs': sum(len(truth[user]) for user in customers),
                  'outsideTruthPairs': sum(len(outside_truth[user]) for user in customers)}
        if counts != EXPECTED_ROLES[role]:
            raise ValueError(f'{role} exact input/truth feasibility differs')
        target = output/f'retail-{role}.csv'
        with target.open('x', encoding='utf-8', newline='') as stream:
            writer = csv.writer(stream, lineterminator='\n')
            writer.writerow(['occurrenceID', 'userID', 'itemID', 'sourceNaiveTime', 'inCatalogue'])
            for index, (customer, stock, _, time) in enumerate(period):
                writer.writerow([f'{role}-{index:06d}', user_map[customer], item_map[stock], time, int(stock in set(catalogue))])
        roles[role] = {'file': target.name, 'sha256': digest(target), 'bytes': target.stat().st_size,
                       'rowsAllItems': len(period), 'counts': counts, 'calendar': [start, end]}
        for customer, stock, _, _ in period:
            past[customer].add(stock)
    manifest = {
        'schema': 'retail-offline-occurrences-v1', 'source': 'UCI Online Retail / Chen 2015',
        'sourceURL': 'https://archive.ics.uci.edu/dataset/352/online+retail',
        'DOI': '10.24432/C5BW33', 'license': 'CC BY 4.0', 'archiveSHA256': ARCHIVE_SHA256,
        'member': member, 'memberSHA256': MEMBER_SHA256, 'derivationSourceSHA256': digest(__file__),
        'filters': filters, 'uniqueThreeMonthSourceOccurrences': len(events),
        'occurrenceUnit': 'unique original(CustomerID,StockCode,InvoiceNo,InvoiceDate); later invoices/times remain repeated events',
        'users': [user_map[raw] for raw in customers], 'items': [item_map[raw] for raw in catalogue],
        'itemOriginalLexicalTieRanks': [lex_order[raw] for raw in catalogue],
        'customerOriginalIDHashes': [hashlib.sha256(raw.encode()).hexdigest() for raw in customers],
        'roles': roles, 'privateIDMapSHA256': digest(output/'private-original-id-map.json'),
        'selection': 'December 100 highest occurrence products/lexical tie; >=3 distinct catalogue products, first 200 canonical CustomerID SHA256/ID tie',
        'limits': ['Derivative subset, not full corpus or anonymity guarantee.',
                   'Counts are invoice-time occurrences, not quantity or price.',
                   'No known exposure, stock availability, dislike or causal conversion label.'],
    }
    write_json_exclusive(output/'retail-manifest.json', manifest)
    return manifest


def load_period(path, expected_sha256, users, catalogue, calendar):
    """No other period file is accessed. Used by the guarded study importer."""
    path = Path(path)
    if digest(path) != expected_sha256:
        raise ValueError('period CSV identity differs from its prepared manifest')
    known_users, known_items = set(users), set(catalogue)
    ids, result = set(), []
    with path.open(encoding='utf-8', newline='') as stream:
        reader = csv.DictReader(stream)
        if reader.fieldnames != ['occurrenceID', 'userID', 'itemID', 'sourceNaiveTime', 'inCatalogue']:
            raise ValueError('offline occurrence schema differs')
        for row in reader:
            if row['occurrenceID'] in ids or row['userID'] not in known_users:
                raise ValueError('duplicate occurrence or unknown selected user ID')
            ids.add(row['occurrenceID'])
            if not calendar[0] <= row['sourceNaiveTime'] < calendar[1]:
                raise ValueError('event is outside its source-naive calendar role')
            if row['inCatalogue'] not in ('0', '1') or (row['itemID'] in known_items) != (row['inCatalogue'] == '1'):
                raise ValueError('local item identity and catalogue marker disagree')
            result.append(row)
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--archive', required=True, type=Path)
    parser.add_argument('--output', required=True, type=Path)
    arguments = parser.parse_args()
    manifest = prepare(arguments.archive, arguments.output)
    print(json.dumps(manifest, indent=2, allow_nan=False))


if __name__ == '__main__':
    main()
