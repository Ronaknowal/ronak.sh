"""One declared Energy geometry-transfer study; preparation never fits a model.

Run --prepare after placing the exact official ZIP in data/source/. Inspect the
derived schema/provenance and pretruth checks before preparing a ready manifest.
--run requires that manifest and an absent exclusive execution marker. It never
refits a winner on validation rows. All JSON is our own numerical state.
"""
import argparse
import csv
import hashlib
import json
import platform
import warnings
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import scipy
import sklearn
from openpyxl import load_workbook
from sklearn.compose import ColumnTransformer, TransformedTargetRegressor
from sklearn.ensemble import RandomForestRegressor
from sklearn.exceptions import ConvergenceWarning
from sklearn.linear_model import MultiTaskLasso, Ridge
from sklearn.multioutput import MultiOutputRegressor
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import OneHotEncoder, StandardScaler

OWNER = Path(__file__).resolve().parent.parent
HEADERS = [f"X{i}" for i in range(1, 9)] + ["Y1", "Y2"]
NUMERIC = [0, 1, 2, 3, 4, 6]
CATEGORICAL = [5, 7]
CATEGORIES = [[2., 3., 4., 5.], [0., 1., 2., 3., 4., 5.]]
SPLIT = {"train": [0, 2, 3, 4, 6, 7, 9, 11], "validation": [1, 8], "test": [5, 10]}
SOURCE_URL = "https://archive.ics.uci.edu/static/public/242/energy+efficiency.zip"
RELEASE_SHA = "596b320893339ed8154b3efbedf0ab7283c2c8d5e74f10d8cb5d7ff9325b7255"


def stamp():
    return datetime.now(timezone.utc).isoformat()


def binding(path):
    raw = Path(path).read_bytes()
    return {"path": Path(path).relative_to(OWNER).as_posix(), "bytes": len(raw),
            "sha256": hashlib.sha256(raw).hexdigest()}


def write_new(path, value):
    """Exclusive evidence creation: an interrupted run cannot erase its history."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("x", encoding="utf-8", newline="\n") as stream:
        json.dump(value, stream, indent=2, allow_nan=False)
        stream.write("\n")


def parse_rows(rows):
    """Permit only fully blank trailing axes; keep original worksheet row IDs."""
    rows = [list(row) for row in rows]
    if not rows:
        raise ValueError("worksheet is empty")
    width = max(map(len, rows))
    rows = [row + [None] * (width - len(row)) for row in rows]
    blank = lambda v: v is None or v == ""
    original_height = len(rows)
    while rows and all(blank(v) for v in rows[-1]):
        rows.pop()
    if not rows:
        raise ValueError("worksheet has no header")
    while width and all(blank(row[width - 1]) for row in rows):
        width -= 1
    if width != 10:
        raise ValueError("require ten nonblank columns X1–X8/Y1–Y2; unexpected columns block preparation")
    if [str(v).strip() for v in rows[0][:10]] != HEADERS:
        raise ValueError("worksheet header must be exactly X1–X8/Y1–Y2")
    if len(rows) - 1 != 768:
        raise ValueError("require exactly 768 nonempty source data rows, with no internal blank row")
    numeric = []
    for row_id, row in enumerate(rows[1:], start=2):
        if any(blank(v) or isinstance(v, bool) for v in row[:10]):
            raise ValueError(f"worksheet row {row_id} has missing/non-numeric fields")
        try:
            numeric.append([float(v) for v in row[:10]])
        except (ValueError, TypeError) as exc:
            raise ValueError(f"worksheet row {row_id} is not a finite numerical row") from exc
    values = np.asarray(numeric)
    if not np.isfinite(values).all():
        raise ValueError("source values must be finite")
    X, Y = values[:, :8], values[:, 8:]
    for column, legal in zip(CATEGORICAL, CATEGORIES):
        if not np.isin(X[:, column], legal).all():
            raise ValueError(f"X{column + 1} contains an undeclared category")
    if any(np.array_equal(X[:, j], Y[:, k]) for j in range(8) for k in range(2)):
        raise ValueError("an input duplicates a target column")
    audit = {"worksheet_rows_including_header": original_height,
             "blank_trailing_rows_discarded": original_height - len(rows),
             "blank_trailing_columns_discarded": len(rows[0]) - width,
             "data_rows": 768, "worksheet_row_ids": list(range(2, 770)),
             "duplicate_input_rows": 768 - len(np.unique(X, axis=0)),
             "duplicate_full_rows": 768 - len(np.unique(values, axis=0)),
             "target_min": Y.min(axis=0).tolist(), "target_max": Y.max(axis=0).tolist()}
    return X, Y, audit


def geometry_split(X):
    X = np.asarray(X, dtype=float)
    if X.shape != (768, 8) or not np.isfinite(X).all():
        raise ValueError("geometry split requires the finite 768×8 source matrix")
    groups, ranks, counts = np.unique(X[:, :5], axis=0, return_inverse=True, return_counts=True)
    if len(groups) != 12 or not np.array_equal(counts, np.full(12, 64)):
        raise ValueError("require twelve lexicographic geometries, each with 64 rows")
    indices = {name: np.flatnonzero(np.isin(ranks, selected)) for name, selected in SPLIT.items()}
    if [len(indices[k]) for k in SPLIT] != [512, 128, 128]:
        raise ValueError("declared geometry split does not have 512/128/128 rows")
    return indices, ranks, groups


def select_data_sheet(workbook):
    """Blank worksheets carry no observations; extra populated sheets fail."""
    populated, empty = [], []
    for sheet in workbook:
        count = sum(any(v is not None and v != "" for v in row)
                    for row in sheet.iter_rows(values_only=True))
        (populated if count else empty).append({"name": sheet.title, "nonblank_rows": count})
    if len(populated) != 1:
        raise ValueError("require exactly one populated data worksheet; extra populated sheets block preparation")
    return workbook[populated[0]["name"]], empty


def prepare():
    archive = OWNER / "data/source/energy-efficiency.zip"
    retrieval_path = OWNER / "data/source/retrieval.json"
    retrieval = json.loads(retrieval_path.read_text(encoding="utf-8"))
    if retrieval["url"] != SOURCE_URL or retrieval["status"] != 200:
        raise ValueError("exact official successful retrieval provenance is required")
    if retrieval["sha256"] != binding(archive)["sha256"]:
        raise ValueError("source ZIP does not match retrieval receipt")
    with zipfile.ZipFile(archive) as packed:
        names = [n for n in packed.namelist() if Path(n).name == "ENB2012_data.xlsx"]
        if len(names) != 1:
            raise ValueError("official archive must contain one ENB2012_data.xlsx")
        raw = packed.read(names[0])
    xlsx = OWNER / "data/source/ENB2012_data.xlsx"
    if xlsx.exists():
        if xlsx.read_bytes() != raw:
            raise ValueError("existing preserved XLSX differs from exact ZIP member")
    else:
        with xlsx.open("xb") as stream:
            stream.write(raw)
    workbook = load_workbook(xlsx, read_only=True, data_only=True)
    try:
        sheet, blank_sheets = select_data_sheet(workbook)
        X, Y, audit = parse_rows(sheet.iter_rows(values_only=True))
        worksheet_name = sheet.title
    finally:
        workbook.close()
    indices, ranks, groups = geometry_split(X)
    std = Y[indices["train"]].std(axis=0, ddof=0)
    if not np.isfinite(std).all() or np.any(std <= 0):
        raise ValueError("nonpositive/nonfinite raw training target scale blocks this study")
    csv_path = OWNER / "data/energy-efficiency.csv"
    with csv_path.open("x", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(HEADERS)
        writer.writerows(np.column_stack([X, Y]).tolist())
    provenance = {"dataset": "UCI Energy Efficiency", "dataset_id": 242,
                  "dataset_doi": "10.24432/C51307", "authors": ["Athanasios Tsanas", "Angeliki Xifara"],
                  "year": 2012, "article_doi": "10.1016/j.enbuild.2012.03.003",
                  "licence": "CC BY 4.0", "licence_url": "https://creativecommons.org/licenses/by/4.0/",
                  "attribution": "Tsanas and Xifara (2012), Energy Efficiency, UCI Machine Learning Repository, DOI10.24432/C51307",
                  "changes": "exact finite source worksheet values serialized to a ten-column CSV; fully blank trailing axes removed with recorded counts; no data-row exclusion",
                  "regime": "computational Ecotect simulation; not field sensor observations",
                  "target_units": "supplied heating/cooling load units", "source_url": SOURCE_URL,
                  "retrieval": retrieval, "source_zip": binding(archive), "source_xlsx": binding(xlsx),
                  "csv": binding(csv_path), "worksheet": worksheet_name,
                  "blank_worksheets": blank_sheets, "headers": HEADERS,
                  "audit": audit, "geometry_tuples": groups.tolist(), "row_geometry_ranks": ranks.tolist(),
                  "split_ranks": SPLIT, "split_csv_row_ids": {k: v.tolist() for k, v in indices.items()},
                  "training_target_mean": Y[indices["train"]].mean(axis=0).tolist(),
                  "training_target_std_ddof0": std.tolist()}
    write_new(OWNER / "data/energy-efficiency-provenance.json", provenance)
    print("Prepared exact source schema / geometry rows / raw training target scales:",
          HEADERS, {k: len(v) for k, v in indices.items()}, std.tolist())
    print("Exact source/CSV/provenance bindings:", binding(archive), binding(xlsx), binding(csv_path),
          binding(OWNER / "data/energy-efficiency-provenance.json"))


def load_data():
    path = OWNER / "data/energy-efficiency.csv"
    with path.open(encoding="utf-8", newline="") as stream:
        rows = list(csv.reader(stream))
    X, Y, audit = parse_rows(rows)
    provenance = json.loads((OWNER / "data/energy-efficiency-provenance.json").read_text(encoding="utf-8"))
    if binding(path) != provenance["csv"]:
        raise ValueError("derived CSV changed after source preparation")
    indices, ranks, groups = geometry_split(X)
    if not np.array_equal(groups, provenance["geometry_tuples"]):
        raise ValueError("geometry identities changed after preparation")
    return X, Y, indices, ranks


def candidate_table():
    rows = []
    for family, settings in [("R", [.1, 1., 10.]), ("S", [1, 4, 12]),
                             ("I", [1, 4, 12]), ("G", [.01, .1, 1.])]:
        for number, value in enumerate(settings, 1):
            params = ({"alpha": value, "fit_intercept": True, "solver": "svd"} if family == "R"
                      else {"alpha": value, "fit_intercept": True, "selection": "cyclic",
                            "tol": 1e-8, "max_iter": 20000} if family == "G"
                      else {"min_samples_leaf": value, "n_estimators": 300, "max_depth": None,
                            "max_features": 1., "bootstrap": True, "min_samples_split": 2,
                            "random_state": 19, "n_jobs": 1})
            rows.append({"id": family + str(number), "family": family, "parameters": params})
    return rows


def estimator(spec):
    family, params = spec["family"], spec["parameters"]
    base = Ridge(**params) if family == "R" else MultiTaskLasso(**params) if family == "G" else RandomForestRegressor(**params)
    if family == "I":
        base = MultiOutputRegressor(base, n_jobs=1)
    encoder = ColumnTransformer([
        ("numeric", "passthrough", NUMERIC),
        ("codes", OneHotEncoder(categories=CATEGORIES, handle_unknown="error", sparse_output=False), CATEGORICAL)
    ], remainder="drop", sparse_threshold=0)
    steps = [("encoder", encoder)]
    if family in ("R", "G"):
        steps.append(("scale", StandardScaler()))
    steps.append(("model", base))
    return TransformedTargetRegressor(regressor=Pipeline(steps), transformer=StandardScaler())


def error_report(truth, prediction, scale, ranks):
    truth, prediction, scale = map(lambda v: np.asarray(v, dtype=float), (truth, prediction, scale))
    ranks = np.asarray(ranks)
    if truth.ndim != 2 or truth.shape != prediction.shape or truth.shape[1] != 2 or len(truth) == 0:
        raise ValueError("error report requires paired nonempty two-output matrices")
    if scale.shape != (2,) or np.any(scale <= 0) or ranks.shape != (len(truth),):
        raise ValueError("error report needs positive training scales and one geometry rank per row")
    if not all(np.isfinite(v).all() for v in (truth, prediction, scale)):
        raise ValueError("nonfinite values cannot become an eligible score")
    residual = prediction - truth
    rmse = np.sqrt(np.mean(residual ** 2, axis=0))
    mae = np.mean(np.abs(residual), axis=0)
    if not np.isfinite(rmse).all() or not np.isfinite(mae).all():
        raise ValueError("error arithmetic overflow")
    groups = {f"G{int(g):02d}": {"rows": int(np.sum(ranks == g)),
              "mean_residual": np.mean(residual[ranks == g], axis=0).tolist(),
              "rmse": np.sqrt(np.mean(residual[ranks == g] ** 2, axis=0)).tolist(),
              "mae": np.mean(np.abs(residual[ranks == g]), axis=0).tolist()}
              for g in np.unique(ranks)}
    return {"rmse": rmse.tolist(), "mae": mae.tolist(),
            "mean_training_std_normalized_rmse": float(np.mean(rmse / scale)), "geometry_errors": groups}


def choose(rows, eligible_ids):
    """Minimum-first absolute tie band; input order is declared table order."""
    valid = [r for r in rows if r["id"] in eligible_ids and r["status"] == "valid"]
    if not valid:
        raise ValueError("selection set has no valid candidate")
    scores = [r["validation"]["mean_training_std_normalized_rmse"] for r in valid]
    if not np.isfinite(scores).all():
        raise ValueError("eligible candidate has a nonfinite validation score")
    minimum = min(scores)
    return next(r["id"] for r, score in zip(valid, scores) if score <= minimum + 1e-12)


def tree_state(tree):
    t = tree.tree_
    return {"children_left": t.children_left.tolist(), "children_right": t.children_right.tolist(),
            "feature": t.feature.tolist(), "threshold": t.threshold.tolist(),
            "values": t.value[:, :, 0].tolist(), "n_outputs": int(tree.n_outputs_)}


def serialize_model(fitted, spec):
    pipe, target = fitted.regressor_, fitted.transformer_
    encoder = pipe.named_steps["encoder"]
    codes = encoder.named_transformers_["codes"]
    state = {"id": spec["id"], "family": spec["family"], "numeric_columns": NUMERIC,
             "categorical_columns": CATEGORICAL, "categories": [v.tolist() for v in codes.categories_],
             "feature_names": encoder.get_feature_names_out().tolist(),
             "target_mean": target.mean_.tolist(), "target_scale": target.scale_.tolist(),
             "target_schema": ["heating", "cooling"], "training_rows": 512}
    if "scale" in pipe.named_steps:
        scale = pipe.named_steps["scale"]
        state["input_mean"], state["input_scale"] = scale.mean_.tolist(), scale.scale_.tolist()
    model = pipe.named_steps["model"]
    if spec["family"] in ("R", "G"):
        state.update({"coef_output_by_feature": model.coef_.tolist(), "intercept": model.intercept_.tolist()})
    else:
        forests = model.estimators_ if spec["family"] == "I" else [model]
        state["forests"] = [{"output_index": i if spec["family"] == "I" else None,
                             "trees": [{"member_id": j, **tree_state(t)} for j, t in enumerate(forest.estimators_)]}
                            for i, forest in enumerate(forests)]
        state["comparison"] = "encoded input cast to IEEE float32 before each <= float64 threshold comparison"
    return state


def predict_saved(state, X):
    """Reconstruct this exact declared preprocessing and fitted numerical state."""
    X = np.asarray(X, dtype=float)
    if X.ndim != 2 or X.shape[1] != 8 or not np.isfinite(X).all():
        raise ValueError("saved-model input requires finite rows with eight original features")
    parts = [X[:, state["numeric_columns"]]]
    for column, categories in zip(state["categorical_columns"], state["categories"]):
        if not np.isin(X[:, column], categories).all():
            raise ValueError("saved-model query contains an unsupported category")
        parts.append(np.column_stack([X[:, column] == code for code in categories]).astype(float))
    encoded = np.column_stack(parts)
    if "input_mean" in state:
        encoded = (encoded - np.asarray(state["input_mean"])) / np.asarray(state["input_scale"])
    if state["family"] in ("R", "G"):
        standardized = encoded @ np.asarray(state["coef_output_by_feature"]).T + np.asarray(state["intercept"])
    else:
        encoded = encoded.astype(np.float32)
        columns = []
        for forest in state["forests"]:
            total = np.zeros((len(X), 1 if forest["output_index"] is not None else 2))
            for tree in forest["trees"]:
                for row, point in enumerate(encoded):
                    node = 0
                    while tree["children_left"][node] != -1:
                        feature = tree["feature"][node]
                        node = (tree["children_left"][node] if float(point[feature]) <= tree["threshold"][node]
                                else tree["children_right"][node])
                    total[row] += tree["values"][node]
            columns.append(total / len(forest["trees"]))
        standardized = np.column_stack(columns)
    prediction = standardized * np.asarray(state["target_scale"]) + np.asarray(state["target_mean"])
    if not np.isfinite(prediction).all():
        raise ValueError("saved prediction arithmetic is nonfinite")
    return prediction


def run(ready_path):
    """Fit every declared candidate once; preserve failures before selection."""
    evidence = OWNER / "content-evidence"
    ready = json.loads(Path(ready_path).read_text(encoding="utf-8"))
    if ready.get("status") != "ready-for-declared-once-study" or ready.get("release_sha256") != RELEASE_SHA:
        raise ValueError("complete pretruth readiness manifest and exact released design required")
    required = [Path(__file__), OWNER / "data/energy-efficiency.csv", OWNER / "data/energy-efficiency-provenance.json",
                OWNER / "reading-flow-design-root-release.json", OWNER / "check-energy-multioutput-pretruth.py",
                OWNER / "content-evidence/energy-multioutput-pretruth-checks.json"]
    if binding(OWNER / "reading-flow-design-root-release.json")["sha256"] != RELEASE_SHA:
        raise ValueError("root release bytes changed")
    for path in required:
        actual = binding(path)
        if actual not in ready["bindings"]:
            raise ValueError(f"readiness does not bind current prerequisite: {actual['path']}")
    if any(evidence.glob("energy-multioutput*execution-marker.json")) or (OWNER / "data/energy-multioutput-results.json").exists():
        raise ValueError("a prior Energy execution marker/result exists; do not replay the campaign")
    X, Y, indices, ranks = load_data()
    train, validation, test = [indices[name] for name in SPLIT]
    scale = Y[train].std(axis=0, ddof=0)
    if not np.isfinite(scale).all() or np.any(scale <= 0):
        raise ValueError("declared study blocks a constant/nonfinite training target")
    specs = candidate_table()
    marker = {"started_utc": stamp(), "ready": binding(Path(ready_path)), "bindings": ready["bindings"],
              "candidate_table": specs, "split_ranks": SPLIT, "no_refit": True,
              "versions": {"python": platform.python_version(), "numpy": np.__version__,
                           "scipy": scipy.__version__, "sklearn": sklearn.__version__}}
    write_new(evidence / "energy-multioutput-execution-marker.json", marker)
    outcomes, fitted_models, predictions = [], {}, {}
    for spec in specs:
        row = {**spec, "status": "failed", "warnings": [], "train": None, "validation": None, "test": None}
        caught = []
        try:
            model = estimator(spec)
            with warnings.catch_warnings(record=True) as caught:
                warnings.simplefilter("always")
                model.fit(X[train], Y[train])
            row["warnings"] = [{"category": type(w.message).__name__, "message": str(w.message)} for w in caught]
            p_train, p_validation = model.predict(X[train]), model.predict(X[validation])
            predictions[spec["id"]] = {"train": p_train.tolist(), "validation": p_validation.tolist()}
            row["train"] = error_report(Y[train], p_train, scale, ranks[train])
            row["validation"] = error_report(Y[validation], p_validation, scale, ranks[validation])
            if spec["family"] == "G":
                core = model.regressor_.named_steps["model"]
                design = model.regressor_[:-1].transform(X[train])
                targets = model.transformer_.transform(Y[train])
                residual = targets - (design @ core.coef_.T + core.intercept_)
                objective = np.sum(residual ** 2) / (2 * len(train)) + spec["parameters"]["alpha"] * np.linalg.norm(core.coef_, axis=0).sum()
                dual_bound = spec["parameters"]["tol"] * np.mean(np.sum((targets - targets.mean(axis=0)) ** 2, axis=1))
                row["convergence"] = {"n_iter": int(core.n_iter_), "dual_gap": float(core.dual_gap_),
                                      "dual_bound": float(dual_bound), "objective": float(objective)}
                if any(issubclass(w.category, ConvergenceWarning) for w in caught) or not np.isfinite(objective) or core.dual_gap_ > dual_bound:
                    raise ValueError("declared shared-row convergence/objective guard failed")
            row["status"] = "valid"
            fitted_models[spec["id"]] = model
        except Exception as exc:
            row["failure"] = {"type": type(exc).__name__, "message": str(exc)}
        finally:
            row["warnings"] = [{"category": type(w.message).__name__, "message": str(w.message)} for w in caught]
        outcomes.append(row)
        write_new(evidence / f"energy-multioutput-candidate-{spec['id']}.json", {"finished_utc": stamp(), "outcome": row,
                  "predictions": predictions.get(spec["id"]), "source": binding(Path(__file__))})
        print("Candidate outcome:", json.dumps(row, allow_nan=False))
    dummy_mean = Y[train].mean(axis=0)
    baseline = {"id": "dummy", "status": "valid", "prediction": dummy_mean.tolist(),
                "train": error_report(Y[train], np.tile(dummy_mean, (len(train), 1)), scale, ranks[train]),
                "validation": error_report(Y[validation], np.tile(dummy_mean, (len(validation), 1)), scale, ranks[validation]), "test": None}
    # No test truth enters error_report, choose, or any fit before this immutable lock.
    try:
        family_winners = {family: choose(outcomes, [s["id"] for s in specs if s["family"] == family]) for family in ("R", "S", "I", "G")}
        overall = choose(outcomes, [s["id"] for s in specs])
        best_independent = choose(outcomes, [s["id"] for s in specs if s["family"] in ("R", "I")])
    except Exception as exc:
        write_new(evidence / "energy-multioutput-blocked.json", {"finished_utc": stamp(), "outcomes": outcomes,
                  "baseline": baseline, "reason": str(exc), "test_scored": False})
        raise
    lock = {"locked_utc": stamp(), "overall_fitted": overall, "families": family_winners,
            "best_independent": best_independent, "baseline_beats_fitted": baseline["validation"]["mean_training_std_normalized_rmse"] < next(r for r in outcomes if r["id"] == overall)["validation"]["mean_training_std_normalized_rmse"],
            "objective": "validation mean_j RMSE_j / raw training_std_j", "tie_tolerance_absolute": 1e-12,
            "tie_rule": "global minimum then first table-order score <= minimum + tolerance",
            "eligible_scores": {r["id"]: r["validation"]["mean_training_std_normalized_rmse"] for r in outcomes if r["status"] == "valid"},
            "test_scored": False, "refit": False, "source": binding(Path(__file__)), "marker": binding(evidence / "energy-multioutput-execution-marker.json")}
    lock_path = OWNER / "data/energy-multioutput-selection-lock.json"
    write_new(lock_path, lock)
    for row in outcomes:
        if row["id"] not in fitted_models:
            continue
        p_test = fitted_models[row["id"]].predict(X[test])
        predictions[row["id"]]["test"] = p_test.tolist()
        row["test"] = error_report(Y[test], p_test, scale, ranks[test])
    baseline["test"] = error_report(Y[test], np.tile(dummy_mean, (len(test), 1)), scale, ranks[test])
    roles = {"overall": overall, "best_independent": best_independent, **{f"family_{k}": v for k, v in family_winners.items()}}
    states, parity = {}, {}
    for model_id in dict.fromkeys(roles.values()):
        spec = next(s for s in specs if s["id"] == model_id)
        states[model_id] = serialize_model(fitted_models[model_id], spec)
        expected = fitted_models[model_id].predict(X)
        rebuilt = predict_saved(states[model_id], X)
        delta = float(np.max(np.abs(rebuilt - expected)))
        parity[model_id] = {"all_source_rows": len(X), "max_abs_error": delta, "passed": bool(np.allclose(rebuilt, expected, rtol=1e-11, atol=1e-10))}
        if not parity[model_id]["passed"]:
            write_new(evidence / "energy-multioutput-reconstruction-failure.json", {"lock": binding(lock_path), "parity": parity})
            raise ValueError("saved-state reconstruction mismatch; preserve campaign and route repair")
    model_path = OWNER / "data/energy-multioutput-models.json"
    write_new(model_path, {"lock": binding(lock_path), "roles": roles, "models": states, "same_frozen_estimator_parity": parity})
    result = {"finished_utc": stamp(), "marker": binding(evidence / "energy-multioutput-execution-marker.json"),
              "source": binding(Path(__file__)), "lock": binding(lock_path), "model_state": binding(model_path),
              "row_ids": {k: v.tolist() for k, v in indices.items()}, "row_geometry_ranks": ranks.tolist(),
              "training_target_std": scale.tolist(), "target_units": "supplied heating/cooling load units",
              "candidate_outcomes": outcomes, "baseline": baseline, "predictions": predictions,
              "source_target_rows": Y.tolist(), "reconstruction": parity, "no_refit": True,
              "scope": "held geometries within supplied simulation; test descriptive after validation lock"}
    result_path = OWNER / "data/energy-multioutput-results.json"
    write_new(result_path, result)
    print("Immutable validation lock:", json.dumps(lock, allow_nan=False))
    print("All descriptive post-lock test outcomes:", json.dumps([{k: r.get(k) for k in ("id", "status", "test", "failure")} for r in outcomes] + [baseline], allow_nan=False))
    print("Same frozen estimator reconstruction:", json.dumps(parity))
    print("Complete source/result/model/lock bindings:", binding(Path(__file__)), binding(result_path), binding(model_path), binding(lock_path))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    modes = parser.add_mutually_exclusive_group(required=True)
    modes.add_argument("--prepare", action="store_true")
    modes.add_argument("--run", type=Path, metavar="PRETRUTH_READY_JSON")
    args = parser.parse_args()
    prepare() if args.prepare else run(args.run)
